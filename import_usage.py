# -*- coding: utf-8 -*-
"""导入「住宿区域用电情况.xlsx」（数据中心导出的每日用电量）。

学校电表余额一天才更新一次，本地记录的余额差估速率会很钝；这份数据直接给出
每天的用电量，导入后估算器用它做种子速率（近 7 天日均 ÷ 24），预测立刻变准。
建议每隔几周从 my.fudan.edu.cn 数据中心导出一次，重新导入刷新。

用法: python import_usage.py [xlsx路径]   （默认为本目录下 住宿区域用电情况.xlsx）
"""

import glob
import os
import sys
from datetime import datetime, timedelta

from openpyxl import load_workbook

from storage import StateStore

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_STATE = os.path.join(BASE_DIR, "state.json")


def parse_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c) for c in rows[0]]
    try:
        i_date = next(i for i, h in enumerate(header) if "日期" in h)
        i_kwh = next(i for i, h in enumerate(header) if "用电" in h)
    except StopIteration:
        raise SystemExit("表头里没找到「日期」和「用电量」列: %s" % header)
    daily = {}
    for r in rows[1:]:
        if r[i_date] is None:
            continue
        d = str(r[i_date])[:10]  # "2026-10-18 00:00:00" -> "2026-10-18"
        datetime.strptime(d, "%Y-%m-%d")  # 校验格式
        daily[d] = daily.get(d, 0.0) + float(r[i_kwh] or 0)
    return sorted(({"date": d, "kwh": k} for d, k in daily.items()),
                  key=lambda x: x["date"])


def summarize(daily, surplus=None):
    now = datetime.now()
    cutoff = now.date() - timedelta(days=6)
    win = [x["kwh"] for x in daily
           if datetime.strptime(x["date"], "%Y-%m-%d").date() >= cutoff]
    span = 7
    if win:
        first = min(x["date"] for x in daily)
        last = max(x["date"] for x in daily)
        d0 = datetime.strptime(first, "%Y-%m-%d").date()
        d1 = datetime.strptime(last, "%Y-%m-%d").date()
        span = min(7, (min(d1, now.date()) - max(d0, cutoff)).days + 1)
    avg = sum(win) / max(1, span)
    print("已导入 %d 天日用电记录（%s ~ %s）" % (
        len(daily), daily[0]["date"] if daily else "-", daily[-1]["date"] if daily else "-"))
    print("近 7 天日均: %.2f 度（%.3f 度/小时）" % (avg, avg / 24))
    if surplus is not None and avg > 0:
        print("按当前剩余 %.1f 度估算: 还能用约 %.1f 小时" % (surplus, surplus / (avg / 24)))


def main():
    args = sys.argv[1:]
    if args:
        path = args[0]
    else:
        cands = [p for p in glob.glob(os.path.join(BASE_DIR, "*.xlsx"))
                 if not os.path.basename(p).startswith("~$")]
        if not cands:
            raise SystemExit("本目录没找到 xlsx，请把「住宿区域用电情况.xlsx」放进本目录，"
                             "或: python import_usage.py <xlsx路径>")
        path = max(cands, key=os.path.getmtime)
    daily = parse_xlsx(path)
    state = StateStore(DEFAULT_STATE)
    surplus = (state.data.get("last_info") or {}).get("surplus")
    state.set_daily_usage(daily)
    state.save()
    print("已写入 state.json（来源: %s）" % os.path.basename(path))
    summarize(daily, surplus)


if __name__ == "__main__":
    main()
